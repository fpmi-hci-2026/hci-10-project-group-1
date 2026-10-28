#!/usr/bin/env python3
"""Сборка отчёта в DOCX напрямую через python-docx.

Конвертация через LibreOffice давала плохую вёрстку: таблицы выходили за поле,
иллюстрации не выравнивались. Здесь документ строится программно: стили абзацев,
таблицы с заданными долями колонок, встроенные изображения, колонтитул с номером
страницы, разрывы перед разделами.

Запуск: .venv/bin/python tools/render_docx.py
"""
from __future__ import annotations
import pathlib, re, sys

ROOT = pathlib.Path(__file__).resolve().parent.parent
MD = ROOT / "ОТЧЁТ" / "ЛабРабота5_ПЧМИ.md"
OUT = ROOT / "ОТЧЁТ" / "ЛабРабота5_ПЧМИ.docx"

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT, MONO = "Times New Roman", "Courier New"
TEXT_W = Cm(16.5)


def set_font(run, size=14, bold=False, italic=False, mono=False):
    run.font.name = MONO if mono else FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run._element.rPr.rFonts.set(qn("w:eastAsia"), MONO if mono else FONT)


def add_runs(par, text: str, size=14):
    """Разбирает **жирный**, *курсив* и `моноширинный` внутри абзаца."""
    for part in re.split(r"(\*\*[^*]+\*\*|(?<![\w*])\*[^*\n]+\*(?![\w*])|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            set_font(par.add_run(part[2:-2]), size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            set_font(par.add_run(part[1:-1]), size - 3, mono=True)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            set_font(par.add_run(part[1:-1]), size, italic=True)
        else:
            set_font(par.add_run(re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", part)), size)


def page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for el, attr in (("w:fldChar", {"w:fldCharType": "begin"}),
                     ("w:instrText", None), ("w:fldChar", {"w:fldCharType": "end"})):
        e = OxmlElement(el)
        if el == "w:instrText":
            e.set(qn("xml:space"), "preserve")
            e.text = " PAGE "
        else:
            for k, v in (attr or {}).items():
                e.set(qn(k), v)
        run._r.append(e)
    set_font(run, 11)


def build() -> int:
    md = MD.read_text(encoding="utf-8")
    doc = Document()

    st = doc.styles["Normal"]
    st.font.name, st.font.size = FONT, Pt(14)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = st.paragraph_format
    pf.line_spacing, pf.space_after, pf.first_line_indent = 1.5, Pt(6), Cm(1.25)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin = Cm(3), Cm(1.5)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
    page_number_footer(sec)

    lines, i, n = md.splitlines(), 0, len(md.splitlines())
    first_h2 = True
    n_tbl = n_fig = 0

    while i < n:
        ln = lines[i]
        if not ln.strip():
            i += 1
            continue

        # титульный лист и содержание размечены сырым HTML — собираем отдельно
        if ln.strip() == "<!--raw-->":
            i += 1
            raw = []
            while i < n and lines[i].strip() != "<!--/raw-->":
                raw.append(lines[i])
                i += 1
            i += 1
            for r in raw:
                m = re.match(r'\s*<p class="(tp-[\w-]+|)"?[^>]*>(.*?)</p>', r) or \
                    re.match(r"\s*<p>(.*?)</p>", r)
                if not m:
                    if 'class="sp-' in r:
                        e = doc.add_paragraph()
                        e.paragraph_format.space_after = Pt(0)
                        set_font(e.add_run(" "), 14)
                    continue
                cls = m.group(1) if m.lastindex == 2 else ""
                txt = re.sub(r"<[^>]+>", "", m.group(m.lastindex))
                par = doc.add_paragraph()
                par.paragraph_format.first_line_indent = Cm(0)
                par.paragraph_format.space_after = Pt(3)
                if cls == "tp-right":
                    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                else:
                    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_font(par.add_run(txt), 19 if cls == "tp-h" else
                         (13 if cls == "tp-org" else 14), bold=(cls == "tp-h"))
            continue

        if ln.strip() == "<!--pagebreak-->":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            i += 1
            continue
        if ln.strip() == "<!--TOC-->":
            par = doc.add_paragraph()
            par.paragraph_format.first_line_indent = Cm(0)
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(par.add_run("Содержание"), 16, bold=True)
            note = doc.add_paragraph()
            note.paragraph_format.first_line_indent = Cm(0)
            set_font(note.add_run(
                "Оглавление с номерами страниц приведено в PDF-версии отчёта."), 12, italic=True)
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            i += 1
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", ln)
        if m:
            lvl, txt = len(m.group(1)), re.sub(r"[*`]", "", m.group(2))
            if lvl == 2 and not first_h2:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            if lvl == 2:
                first_h2 = False
            par = doc.add_paragraph()
            par.paragraph_format.first_line_indent = Cm(0)
            par.paragraph_format.space_before = Pt(10 if lvl > 2 else 0)
            par.paragraph_format.space_after = Pt(8)
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_font(par.add_run(txt), {1: 18, 2: 16, 3: 14, 4: 14}[lvl],
                     bold=True, italic=(lvl == 4))
            i += 1
            continue

        mi = re.match(r"^!\[([^\]]*)\]\(([^)\s]+)\)\s*$", ln)
        if mi:
            img = MD.parent / mi.group(2)
            if img.exists():
                n_fig += 1
                par = doc.add_paragraph()
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.paragraph_format.first_line_indent = Cm(0)
                par.add_run().add_picture(str(img), width=TEXT_W)
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.paragraph_format.first_line_indent = Cm(0)
                set_font(cap.add_run(re.sub(r"[*`]", "", mi.group(1))), 12)
            i += 1
            continue

        if ln.lstrip().startswith("|") and i + 1 < n and re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i + 1]):
            head = [c.strip() for c in ln.strip().strip("|").split("|")]
            i += 2
            rows = []
            while i < n and lines[i].lstrip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            n_tbl += 1
            t = doc.add_table(rows=1, cols=len(head))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = False
            avg = [max(4, sum([len(head[k])] + [len(r[k]) for r in rows if k < len(r)]) /
                       (1 + len([r for r in rows if k < len(r)]))) for k in range(len(head))]
            share = [min(0.42, max(0.07, v / sum(avg))) for v in avg]
            share = [x / sum(share) for x in share]
            for k, c in enumerate(t.rows[0].cells):
                c.width = Cm(16.5 * share[k])
                p0 = c.paragraphs[0]
                p0.paragraph_format.first_line_indent = Cm(0)
                p0.paragraph_format.space_after = Pt(0)
                p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
                set_font(p0.add_run(re.sub(r"[*`]", "", head[k])), 10.5, bold=True)
            for r in rows:
                cells = t.add_row().cells
                for k, c in enumerate(cells):
                    if k >= len(r):
                        continue
                    c.width = Cm(16.5 * share[k])
                    p0 = c.paragraphs[0]
                    p0.paragraph_format.first_line_indent = Cm(0)
                    p0.paragraph_format.space_after = Pt(0)
                    p0.paragraph_format.line_spacing = 1.15
                    p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    add_runs(p0, r[k], 10.5)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        if ln.lstrip().startswith(">"):
            buf = []
            while i < n and lines[i].lstrip().startswith(">"):
                buf.append(lines[i].lstrip()[1:].strip())
                i += 1
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(1.25)
            par.paragraph_format.first_line_indent = Cm(0)
            add_runs(par, " ".join(buf))
            for r in par.runs:
                r.italic = True
            continue

        for pat, style in ((r"^\s*[*-]\s+", "List Bullet"), (r"^\s*\d+\.\s+", "List Number")):
            if re.match(pat, ln):
                while i < n and re.match(pat, lines[i]):
                    par = doc.add_paragraph(style=style)
                    par.paragraph_format.first_line_indent = Cm(0)
                    par.paragraph_format.left_indent = Cm(1.25)
                    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    add_runs(par, re.sub(pat, "", lines[i]))
                    i += 1
                break
        else:
            buf = []
            while i < n and lines[i].strip() and not re.match(
                    r"^(#{1,4}\s|!\[|<!--|\s*[|>*-]|\s*\d+\.\s)", lines[i]):
                buf.append(lines[i].strip())
                i += 1
            if not buf:
                i += 1
                continue
            text = " ".join(buf)
            par = doc.add_paragraph()
            if re.match(r"^\*\*Таблица \d+", text):
                par.paragraph_format.first_line_indent = Cm(0)
                par.paragraph_format.space_before = Pt(10)
            add_runs(par, text)

    doc.save(str(OUT))
    print(f"  готово: {OUT.name} — таблиц {n_tbl}, иллюстраций {n_fig}, "
          f"{OUT.stat().st_size // 1024} КБ")
    return 0


if __name__ == "__main__":
    sys.exit(build())
